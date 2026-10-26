"""add_refs_46_50.py - append references [46]-[50] before 'Supplementary material'
heading, cloning the [45] Bibliography paragraph format."""
import copy
import re
import docx
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

P = r"D:\AL_FuzzyClustering_Paper\v2_unified\draft\DRAL_revised_final.docx"
d = docx.Document(P)
ps = d.paragraphs

# 1. verify in-text citations exist
alltext = "\n".join(p.text for p in ps)
for n in range(46, 51):
    hits = len(re.findall(rf"\[{n}\]", alltext))
    print(f"in-text [{n}] occurrences: {hits}")

# 2. locate anchor: the 'Supplementary material' Heading 1 (first occurrence after refs)
anchor = None
for p in ps:
    if p.style.name == "Heading 1" and "Supplementary" in p.text:
        anchor = p
        break
assert anchor is not None, "anchor heading not found"
print("anchor found at:", anchor.text)

# 3. clone [45] format from its paragraph
tpl = None
for p in ps:
    if p.text.strip().startswith("[45]"):
        tpl = p
        break
assert tpl is not None

refs = [
    "[46] A.R. Jonckheere, 1954. A distribution-free k-sample test against ordered alternatives. Biometrika, 41(1-2), 133-145. https://doi.org/10.1093/biomet/41.1-2.133",
    "[47] R. DerSimonian and N. Laird, 1986. Meta-analysis in clinical trials. Controlled Clinical Trials, 7(3), 177-188. https://doi.org/10.1016/0197-2456(86)90046-2",
    "[48] S. Holm, 1979. A simple sequentially rejective multiple test procedure. Scandinavian Journal of Statistics, 6(2), 65-70. https://www.jstor.org/stable/4615733",
    "[49] Y. Benjamini and Y. Hochberg, 1995. Controlling the false discovery rate: a practical and powerful approach to multiple testing. Journal of the Royal Statistical Society. Series B (Methodological), 57(1), 289-300. https://doi.org/10.1111/j.2517-6161.1995.tb02031.x",
    "[50] J. Cohen, 1988. Statistical Power Analysis for the Behavioral Sciences (2nd ed.). Lawrence Erlbaum Associates. https://doi.org/10.4324/9780203771587",
]

def clone_bib_before(anchor_p, template_p, text):
    new_el = copy.deepcopy(template_p._p)
    for r in new_el.findall(qn("w:r")):
        new_el.remove(r)
    anchor_p._p.addprevious(new_el)
    para = Paragraph(new_el, anchor_p._parent)
    para.add_run(text)
    return para

for t in refs:
    clone_bib_before(anchor, tpl, t)
    print("inserted:", t[:60])

d.save(P)
# 4. verify
d2 = docx.Document(P)
bib = [p.text for p in d2.paragraphs if p.text.strip().startswith("[")]
print("total bibliography entries now:", len(bib))
print("last 5:", [b[:60] for b in bib[-5:]])
