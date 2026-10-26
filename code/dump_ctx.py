"""dump_ctx.py - print full text of citation-relevant paragraphs with context."""
import docx

P = r"D:\AL_FuzzyClustering_Paper\v2_unified\draft\DRAL_revised_final.docx"
d = docx.Document(P)
wanted = list(range(193, 226)) + [453]
for i in wanted:
    if i >= len(d.paragraphs):
        continue
    t = d.paragraphs[i].text
    print(f"--- {i} [{d.paragraphs[i].style.name}] ---")
    print(t)
    print()
