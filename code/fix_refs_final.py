"""fix_refs_final.py - replace any empty Bibliography paragraphs at the tail of
the References block with properly built [46]-[50] entries (lxml run building)."""
import copy
import docx
from docx.oxml.ns import qn

P = r"D:\AL_FuzzyClustering_Paper\v2_unified\draft\DRAL_revised_final.docx"
d = docx.Document(P)
ps = d.paragraphs

# locate References block boundaries
refs_h = [i for i, p in enumerate(ps) if p.style.name == "Heading 1" and p.text.strip() == "References"][0]
sup_h = [i for i, p in enumerate(ps) if p.style.name == "Heading 1" and p.text.strip() == "Supplementary material"][0]
print("refs block:", refs_h, "->", sup_h)

# tail Bibliography paragraphs after [45] (empty or [46..50])
tpl = None
empty_tail = []
for i in range(sup_h - 1, refs_h, -1):
    p = ps[i]
    t = p.text.strip()
    if t.startswith("[45]"):
        tpl = p
        break
    if p.style.name == "Bibliography":
        empty_tail.append(p)
print("template [45] at:", next(i for i, p in enumerate(ps) if p.text.strip().startswith("[45]")))
print("tail empty bib paras:", len(empty_tail))

# remove stale tail paragraphs (regardless of content, we rebuild [46..50])
for p in empty_tail:
    p._p.getparent().remove(p._p)

refs_text = [
    "[46] A.R. Jonckheere, 1954. A distribution-free k-sample test against ordered alternatives. Biometrika, 41(1-2), 133-145. https://doi.org/10.1093/biomet/41.1-2.133",
    "[47] R. DerSimonian and N. Laird, 1986. Meta-analysis in clinical trials. Controlled Clinical Trials, 7(3), 177-188. https://doi.org/10.1016/0197-2456(86)90046-2",
    "[48] S. Holm, 1979. A simple sequentially rejective multiple test procedure. Scandinavian Journal of Statistics, 6(2), 65-70. https://www.jstor.org/stable/4615733",
    "[49] Y. Benjamini and Y. Hochberg, 1995. Controlling the false discovery rate: a practical and powerful approach to multiple testing. Journal of the Royal Statistical Society. Series B (Methodological), 57(1), 289-300. https://doi.org/10.1111/j.2517-6161.1995.tb02031.x",
    "[50] J. Cohen, 1988. Statistical Power Analysis for the Behavioral Sciences (2nd ed.). Lawrence Erlbaum Associates. https://doi.org/10.4324/9780203771587",
]

def build_bib_paragraph(template_p, text):
    new_el = copy.deepcopy(template_p._p)
    # drop everything except w:pPr
    for child in list(new_el):
        if child.tag != qn("w:pPr"):
            new_el.remove(child)
    r = new_el.makeelement(qn("w:r"), {})
    t = new_el.makeelement(qn("w:t"), {})
    t.text = text
    r.append(t)
    new_el.append(r)
    return new_el

# anchor again (element may have shifted after removals)
d2_anchor = None
for p in d.paragraphs:
    if p.style.name == "Heading 1" and p.text.strip() == "Supplementary material":
        d2_anchor = p
        break
for text in refs_text:
    el = build_bib_paragraph(tpl, text)
    d2_anchor._p.addprevious(el)
    print("inserted:", text[:50])

d.save(P)

# verify
d3 = docx.Document(P)
ps3 = d3.paragraphs
refs3 = [i for i, p in enumerate(ps3) if p.style.name == "Heading 1" and p.text.strip() == "References"][0]
sup3 = [i for i, p in enumerate(ps3) if p.style.name == "Heading 1" and p.text.strip() == "Supplementary material"][0]
bib_tail = [(i, ps3[i].text.strip()) for i in range(sup3 - 7, sup3)]
for i, t in bib_tail:
    print(i, t[:80] if t else "<EMPTY>")
bib_before = sum(1 for i, p in enumerate(ps3) if p.style.name == "Bibliography" and i < refs3)
bib_total = sum(1 for p in ps3 if p.style.name == "Bibliography")
print("bib before References:", bib_before, "| total bib:", bib_total)
